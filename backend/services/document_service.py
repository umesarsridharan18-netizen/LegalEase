from io import BytesIO
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfbase import pdfmetrics
from reportlab.lib import colors

from backend.utils.text_utils import sanitize_text

BASE_DIR = Path(__file__).resolve().parents[2]
LOGO_PATH = BASE_DIR / "assets" / "logo.png"


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def _add_docx_logo(doc: Document) -> None:
    if LOGO_PATH.exists():
        section = doc.sections[0]
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(str(LOGO_PATH), width=Inches(1.35))


def format_docx(text: str, doc_type: str) -> bytes:
    text = sanitize_text(text)
    doc = Document()
    _add_docx_logo(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    doc.add_paragraph("")

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            doc.add_paragraph("")
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)

        if re.match(r"^(ARTICLE|SECTION|[0-9]+\.)\b", line, re.I):
            r = p.add_run(line)
            r.bold = True
        else:
            p.add_run(line)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase — AI-assisted legal document draft")

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def _register_font() -> str:
    candidates = [
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/ARIAL.TTF"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            try:
                pdfmetrics.registerFont(TTFont("LegalEaseSans", str(candidate)))
                return "LegalEaseSans"
            except Exception:
                pass
    return "Helvetica"


def _pdf_header_footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.grey)
    canvas.drawCentredString(
        A4[0] / 2,
        10 * mm,
        "LegalEase — AI-assisted legal document draft"
    )
    canvas.restoreState()


def format_pdf(text: str, doc_type: str) -> bytes:
    text = sanitize_text(text)
    output = BytesIO()
    font = _register_font()

    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=22 * mm,
        bottomMargin=18 * mm,
        title=doc_type,
        author="LegalEase",
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "LegalEaseBody",
        parent=styles["BodyText"],
        fontName=font,
        fontSize=10.5,
        leading=15,
        spaceAfter=8,
    )
    heading = ParagraphStyle(
        "LegalEaseHeading",
        parent=body,
        fontName=font,
        fontSize=11.5,
        leading=15,
        spaceBefore=8,
        spaceAfter=7,
    )
    title_style = ParagraphStyle(
        "LegalEaseTitle",
        parent=body,
        fontName=font,
        fontSize=16,
        leading=20,
        alignment=TA_CENTER,
        spaceAfter=14,
    )

    story = [
        Paragraph("LegalEase", title_style),
        Paragraph(doc_type.upper(), title_style),
    ]

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 4))
            continue

        safe = (
            line.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        if re.match(r"^(ARTICLE|SECTION|[0-9]+\.)\b", line, re.I):
            story.append(Paragraph(f"<b>{safe}</b>", heading))
        else:
            story.append(Paragraph(safe, body))

    doc.build(story, onFirstPage=_pdf_header_footer, onLaterPages=_pdf_header_footer)
    return output.getvalue()
